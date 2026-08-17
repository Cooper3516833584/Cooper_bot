
# handinsvc.py
from __future__ import annotations
import asyncio
import hashlib
import json
import re
import time
import urllib.request
import urllib.error
import urllib.parse
import shutil
from dataclasses import dataclass, asdict, field
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Set
from urllib.parse import urlparse, unquote
import os

from cooper_bot.core.config import (
    DATA_DIR,
    GROUP_DOCS_DIR,
    USER_DOCS_DIR,
    HANDIN_DB_PATH,
    HANDIN_INBOX_DIR,
    HANDIN_TASKS_DIRNAME,
    HANDIN_ROOT_DIR,
    ROSTER_XLSX_PATH,
    TIMEZONE,
    NAPCAT_TEMP_CONTAINER_DIR,
    NAPCAT_TEMP_HOST_DIR,
    HANDIN_KEEP_DAYS_AFTER_LAST_GET,
    HANDIN_INBOX_KEEP_DAYS,
)
from cooper_bot.core.logger import Logger
from cooper_bot.core.ziputil import open_fast_zip, write_path as zip_write_path

try:
    from zoneinfo import ZoneInfo
except Exception:  # Py<3.9
    ZoneInfo = None

import openpyxl


HANDIN_ALLOWED_REQUIRED_SUFFIXES = {
    "7z",
    "bmp",
    "c",
    "cpp",
    "csv",
    "css",
    "doc",
    "docx",
    "gif",
    "heic",
    "html",
    "ipynb",
    "java",
    "jpeg",
    "jpg",
    "js",
    "json",
    "md",
    "pdf",
    "png",
    "ppt",
    "pptx",
    "py",
    "rar",
    "txt",
    "webp",
    "xls",
    "xlsx",
    "xml",
    "zip",
}


def normalize_required_suffix(suffix: str) -> str:
    s = str(suffix or "").strip().lstrip(".").casefold()
    return s if s in HANDIN_ALLOWED_REQUIRED_SUFFIXES else ""


def required_suffix_display(suffix: str) -> str:
    s = normalize_required_suffix(suffix)
    return f".{s}" if s else ""


def file_matches_required_suffix(filename: str, suffix: str) -> bool:
    s = normalize_required_suffix(suffix)
    if not s:
        return True
    return Path(str(filename or "")).suffix.casefold() == f".{s}"


# ========= 文件名提取（参考 who_has_handed_in.py 的逻辑） =========
BLACKLIST_SUBSTRINGS = {
    "电气", "学院", "工程", "班", "专业",
    "报告", "读书", "作业", "论文", "马原",
    "课", "阅读", "历史", "自由", "之间",
    "政治", "经济", "序言", "导言", "经典", "思想","电测","报告","电路测试","测试",
}
STRUCTURAL_WORDS = ["电气", "学院", "工程", "班", "专业"]
SEPARATORS = ["-", "_", "——", "—", "–", ";", "，", ",", " "]

_RE_STU = re.compile(r"[Uu]\d{8,12}")
_RE_ENG = re.compile(r"[A-Za-z]")
_RE_NUM = re.compile(r"[Uu]?\d{4,}")
SUBMITTED_FILE_SUFFIXES = {".doc", ".docx", ".pdf", ".txt", ".zip", ".rar", ".7z", ".ppt", ".pptx", ".xls", ".xlsx"}
HANDIN_HASH_INDEX_FILENAME = ".handin_file_hashes.json"

def clean_filename(filename: str) -> str:
    stem = Path(filename).stem
    for sep in SEPARATORS:
        stem = stem.replace(sep, " ")
    stem = _RE_NUM.sub(" ", stem)
    stem = _RE_ENG.sub(" ", stem)
    return stem

def looks_like_name(token: str) -> bool:
    if not re.fullmatch(r"[\u4e00-\u9fff]{2,3}", token):
        return False
    for bad in BLACKLIST_SUBSTRINGS:
        if bad in token:
            return False
    return True

def extract_name_from_filename(filename: str) -> str:
    part = clean_filename(filename)
    tokens = [t for t in part.split() if t]
    for tok in reversed(tokens):
        if looks_like_name(tok):
            return tok

    for tok in tokens:
        if not re.fullmatch(r"[\u4e00-\u9fff]+", tok):
            continue
        for sw in STRUCTURAL_WORDS:
            idx = tok.find(sw)
            if idx >= 2:
                prefix = tok[:idx]
                if looks_like_name(prefix):
                    return prefix

    chunks = re.findall(r"[\u4e00-\u9fff]+", part)
    candidates = []
    for chunk in chunks:
        for n in (3, 2):
            for i in range(len(chunk) - n + 1):
                sub = chunk[i:i+n]
                if looks_like_name(sub):
                    candidates.append(sub)
    return candidates[-1] if candidates else ""

def extract_student_id(filename: str) -> str:
    m = _RE_STU.search(filename or "")
    return m.group(0).upper() if m else ""


# ========= 名册读取 =========
def load_roster(path: Path) -> List[Tuple[str, str]]:
    """读取班级名册，返回 [(学号, 姓名), ...]。支持首行是标题、第二行才是表头。"""
    path = Path(path)
    if not path.exists():
        return []
    wb = openpyxl.load_workbook(path)
    ws = wb.active

    header_row = None
    col_id = None
    col_name = None

    # 找包含“学号”“姓名”的表头行
    for r in range(1, min(30, ws.max_row) + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, min(50, ws.max_column) + 1)]
        # 统一成字符串比较
        row_str = [str(v).strip() for v in row_vals if v is not None]
        if "学号" in row_str and "姓名" in row_str:
            header_row = r
            # 找列号
            for c in range(1, ws.max_column + 1):
                v = ws.cell(r, c).value
                if v is None:
                    continue
                s = str(v).strip()
                if s == "学号":
                    col_id = c
                elif s == "姓名":
                    col_name = c
            break

    if not header_row or not col_id or not col_name:
        return []

    out: List[Tuple[str, str]] = []
    for r in range(header_row + 1, ws.max_row + 1):
        sid = ws.cell(r, col_id).value
        name = ws.cell(r, col_name).value
        if sid is None and name is None:
            continue
        sid_s = str(sid).strip() if sid is not None else ""
        name_s = str(name).strip() if name is not None else ""
        if not sid_s and not name_s:
            continue
        out.append((sid_s.upper(), name_s))
    return out


# ========= 时间解析 =========
_RE_TIME = re.compile(r"^\s*(\d{1,2})[.\u3002/\-](\d{1,2})\s*(\d{1,2})[:：](\d{1,2})\s*$")

def parse_mmdd_hhmm(s: str, now_ts: float) -> Optional[float]:
    """把 'M.D HH:MM' 解析为时间戳。若解析出的时间 <= 当前，则自动 +1 年。"""
    s = (s or "").strip()
    m = _RE_TIME.match(s)
    if not m:
        return None
    mon = int(m.group(1))
    day = int(m.group(2))
    hh = int(m.group(3))
    mm = int(m.group(4))

    tz = None
    if ZoneInfo:
        try:
            tz = ZoneInfo(TIMEZONE)
        except Exception:
            tz = None

    now = time.time() if now_ts is None else float(now_ts)
    if tz:
        now_dt = time.localtime(now)
        # 用 tz-aware datetime
        import datetime as _dt
        n = _dt.datetime.fromtimestamp(now, tz)
        year = n.year
        dt = _dt.datetime(year, mon, day, hh, mm, tzinfo=tz)
        if dt.timestamp() <= now:
            dt = _dt.datetime(year + 1, mon, day, hh, mm, tzinfo=tz)
        return dt.timestamp()
    else:
        # 退化：用本地时区
        lt = time.localtime(now)
        year = lt.tm_year
        import datetime as _dt
        dt = _dt.datetime(year, mon, day, hh, mm)
        ts = dt.timestamp()
        if ts <= now:
            dt2 = _dt.datetime(year + 1, mon, day, hh, mm)
            ts = dt2.timestamp()
        return ts


def pretty_ts(ts: float) -> str:
    try:
        return time.strftime("%Y-%m-%d %H:%M", time.localtime(ts))
    except Exception:
        return str(ts)


# ========= 任务数据结构 =========
@dataclass
class HandinTask:
    task_id: str
    group_id: int
    creator_id: int
    name: str
    created_ts: float
    required_suffix: str = ""
    # 可选的多个提醒时间（时间戳列表）。最后一个时间一定是截止时间，提醒时间可为空。
    remind_ts_list: List[float] = field(default_factory=list)
    # 已发送到第几个提醒（下一个将发送的提醒索引）
    remind_sent_idx: int = 0
    deadline_ts: float = 0.0
    deadline_sent: bool = False
    closed: bool = False
    cancelled: bool = False
    cancelled_ts: float = 0.0
    cancelled_by: int = 0

    # 任务创建者最后一次 /handinget 的时间戳（用于归档保留策略）
    last_handinget_ts: float = 0.0
    # 归档是否已被清理（清理后 /handinget 不再可用，但日志仍保留）
    purged: bool = False
    purged_ts: float = 0.0

    def is_active(self, now: Optional[float] = None) -> bool:
        now = time.time() if now is None else float(now)
        return (not self.closed) and now < float(self.deadline_ts)


class HandinService:
    """提交任务服务：任务管理 + 提交文件归档 + 未交名单统计 + 定时提醒/截止推送。"""

    def __init__(self, log: Logger):
        self.log = log
        self.db_path = Path(HANDIN_DB_PATH)
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self.inbox_dir = Path(HANDIN_INBOX_DIR)
        self.inbox_dir.mkdir(parents=True, exist_ok=True)

        # 新版：提交文件不再放在 data/groups/<gid>/handin 下，避免群成员通过 /find 看到他人提交
        # 统一放在 data/handin/<gid>/<task>/files/
        self.handin_root = Path(HANDIN_ROOT_DIR)
        self.handin_root.mkdir(parents=True, exist_ok=True)

        # 兼容迁移：把旧版 data/groups/<gid>/handin/<task>/... 搬到 data/handin/<gid>/<task>/...
        self._migrate_legacy_tree()

        self._tasks: Dict[str, HandinTask] = {}
        self._load()

        # 清理节流：避免每 10 秒全盘扫描
        self._last_cleanup_ts: float = 0.0
        # 名册缓存（按 mtime 刷新）
        self._roster_cache_mtime: float = -1.0
        self._roster_cache: List[Tuple[str, str]] = []

    def is_task_gettable(self, task: HandinTask) -> bool:
        """任务是否仍可 /handinget：归档未被清理且目录仍在。"""
        try:
            if getattr(task, "purged", False):
                return False
            files_dir = self._task_files_dir(task.group_id, task.name)
            return files_dir.exists()
        except Exception:
            return False

    def _purge_task_archive(self, task: HandinTask, now: Optional[float] = None) -> bool:
        """删除某任务的归档目录，并标记为 purged。返回是否有变更。"""
        now = time.time() if now is None else float(now)
        try:
            tdir = self._task_dir(task.group_id, task.name)
            if tdir.exists():
                shutil.rmtree(tdir, ignore_errors=True)
        except Exception as e:
            self.log.warning(
                f"[handin] purge archive dir failed: task={task.task_id} group={task.group_id} "
                f"name={task.name} err={e}"
            )
        changed = False
        if not getattr(task, "purged", False):
            task.purged = True
            task.purged_ts = now
            changed = True
        return changed

    def cleanup_archives_and_inbox(self, now: Optional[float] = None) -> bool:
        """清理：
        - 归档：在任务创建者最后一次 /handinget 后保留 HANDIN_KEEP_DAYS_AFTER_LAST_GET 天
        - inbox：临时收件箱内文件保留 HANDIN_INBOX_KEEP_DAYS 天
        返回：是否发生了 DB 变更（需要保存）。
        """
        now = time.time() if now is None else float(now)
        changed = False

        keep_sec = float(HANDIN_KEEP_DAYS_AFTER_LAST_GET) * 86400.0
        # 1) 归档清理
        for t in list(self._tasks.values()):
            if getattr(t, "purged", False):
                continue
            last_get = float(getattr(t, "last_handinget_ts", 0.0) or 0.0)
            if last_get <= 0:
                continue
            # 仍在进行中的任务不清理
            if t.is_active(now):
                continue
            if now - last_get >= keep_sec:
                if self._purge_task_archive(t, now=now):
                    changed = True

        # 2) inbox 清理（按文件 mtime）
        inbox_keep = float(HANDIN_INBOX_KEEP_DAYS) * 86400.0
        try:
            if self.inbox_dir.exists():
                for p in self.inbox_dir.rglob("*"):
                    if not p.is_file():
                        continue
                    try:
                        if now - float(p.stat().st_mtime) >= inbox_keep:
                            p.unlink(missing_ok=True)
                    except Exception as e:
                        self.log.warning(f"[handin] inbox cleanup item failed: path={p} err={e}")
                        continue
        except Exception as e:
            self.log.warning(f"[handin] inbox cleanup failed: root={self.inbox_dir} err={e}")

        return changed

    # ----- persistence -----
    def _load(self):
        try:
            if self.db_path.exists():
                obj = json.loads(self.db_path.read_text(encoding="utf-8"))
                if isinstance(obj, dict):
                    for tid, t in obj.items():
                        if not isinstance(t, dict):
                            continue
                        td = dict(t)

                        # 兼容旧字段：remind_ts/remind_sent
                        if "remind_ts_list" not in td:
                            r = td.get("remind_ts", None)
                            td["remind_ts_list"] = [float(r)] if r is not None else []
                            if td.get("remind_sent") is True and td["remind_ts_list"]:
                                td["remind_sent_idx"] = len(td["remind_ts_list"])
                            else:
                                td["remind_sent_idx"] = 0
                            td.pop("remind_ts", None)
                            td.pop("remind_sent", None)
                        else:
                            td["remind_ts_list"] = [float(x) for x in (td.get("remind_ts_list") or [])]
                            td.setdefault("remind_sent_idx", 0)

                        td.setdefault("deadline_sent", False)
                        td.setdefault("closed", False)
                        td.setdefault("cancelled", False)
                        td.setdefault("cancelled_ts", 0.0)
                        td.setdefault("cancelled_by", 0)
                        td.setdefault("last_handinget_ts", 0.0)
                        td.setdefault("purged", False)
                        td.setdefault("purged_ts", 0.0)
                        td["required_suffix"] = normalize_required_suffix(td.get("required_suffix", ""))
                        self._tasks[str(tid)] = HandinTask(**td)
        except Exception as e:
            self.log.warning(f"Handin DB load failed: {e}")
            self._tasks = {}

    def _save(self):
        try:
            obj = {tid: asdict(t) for tid, t in self._tasks.items()}
            tmp = self.db_path.with_suffix(self.db_path.suffix + ".tmp")
            tmp.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")
            tmp.replace(self.db_path)
        except Exception as e:
            self.log.warning(f"Handin DB save failed: {e}")

    def _get_roster(self) -> List[Tuple[str, str]]:
        """读取并缓存名册（文件 mtime 变化时自动刷新）。"""
        path = Path(ROSTER_XLSX_PATH)
        if not path.exists():
            self._roster_cache = []
            self._roster_cache_mtime = -1.0
            return []
        try:
            mtime = float(path.stat().st_mtime)
        except Exception:
            mtime = -1.0
        if mtime >= 0 and abs(mtime - float(self._roster_cache_mtime)) < 1e-6:
            return list(self._roster_cache)
        try:
            data = load_roster(path)
        except Exception:
            data = []
        self._roster_cache = list(data or [])
        self._roster_cache_mtime = mtime
        return list(self._roster_cache)

    def _get_roster_names(self) -> List[str]:
        names: List[str] = []
        seen: Set[str] = set()
        for _, nm in self._get_roster():
            name = str(nm or "").strip()
            if (not name) or (name in seen):
                continue
            seen.add(name)
            names.append(name)
        names.sort(key=lambda s: len(s), reverse=True)
        return names

    def find_roster_name_in_filename(self, filename: str, roster_names: Optional[List[str]] = None) -> str:
        """在文件名中查找是否包含名册中的姓名，返回首个命中的姓名。"""
        fn = str(filename or "")
        if not fn:
            return ""
        stem = Path(fn).stem
        compact = re.sub(r"\s+", "", stem)
        names = roster_names if roster_names is not None else self._get_roster_names()
        for nm in names:
            if nm and (nm in stem or nm in compact):
                return nm
        return ""

    # ----- paths -----
    def _task_dir(self, group_id: int, task_name: str) -> Path:
        safe = self._safe_component(task_name)
        return (self.handin_root / str(group_id) / safe)

    def _task_files_dir(self, group_id: int, task_name: str) -> Path:
        d = self._task_dir(group_id, task_name) / "files"
        d.mkdir(parents=True, exist_ok=True)
        return d

    def _legacy_handin_dir(self, group_id: int) -> Path:
        """旧版本：data/groups/<gid>/handin"""
        return GROUP_DOCS_DIR / str(group_id) / HANDIN_TASKS_DIRNAME

    @staticmethod
    def _move_or_merge_dir(src: Path, dst: Path, log: Logger) -> None:
        """把 src 目录迁移到 dst。dst 不存在则直接 move；存在则合并并对冲突文件重命名。"""
        src = Path(src)
        dst = Path(dst)
        if not src.exists() or not src.is_dir():
            return
        dst.parent.mkdir(parents=True, exist_ok=True)

        if not dst.exists():
            try:
                shutil.move(str(src), str(dst))
            except Exception as e:
                log.warning(f"[handin] move legacy dir failed: {src} -> {dst}: {e}")
            return

        # dst 已存在：递归合并
        def merge_dir(a: Path, b: Path):
            b.mkdir(parents=True, exist_ok=True)
            for item in a.iterdir():
                target = b / item.name
                if item.is_dir():
                    merge_dir(item, target)
                    try:
                        item.rmdir()
                    except Exception:
                        pass
                else:
                    if target.exists():
                        stem, suf = target.stem, target.suffix
                        for i in range(1, 999):
                            alt = b / f"{stem}_legacy{i}{suf}"
                            if not alt.exists():
                                target = alt
                                break
                    try:
                        shutil.move(str(item), str(target))
                    except Exception as e:
                        log.warning(f"[handin] move legacy file failed: {item} -> {target}: {e}")

        try:
            merge_dir(src, dst)
            # 尝试删除残余
            try:
                shutil.rmtree(src)
            except Exception:
                pass
        except Exception as e:
            log.warning(f"[handin] merge legacy dir failed: {src} -> {dst}: {e}")

    def _migrate_legacy_tree(self) -> None:
        """启动时迁移所有群的旧 handin 目录到 data/handin 下，并尽量清理旧目录。"""
        try:
            if not GROUP_DOCS_DIR.exists():
                return

            moved_any = False
            for gdir in GROUP_DOCS_DIR.iterdir():
                if not gdir.is_dir():
                    continue
                gid = gdir.name
                legacy = gdir / HANDIN_TASKS_DIRNAME
                if not legacy.exists() or not legacy.is_dir():
                    continue

                dst_gid = self.handin_root / gid
                for task_dir in legacy.iterdir():
                    if not task_dir.is_dir():
                        continue
                    dst_task = dst_gid / task_dir.name
                    self._move_or_merge_dir(task_dir, dst_task, self.log)
                    moved_any = True

                # 尝试删除旧 handin 目录（空则成功）
                try:
                    legacy.rmdir()
                except Exception:
                    pass

            if moved_any:
                self.log.info("[handin] migrated legacy submissions into data/handin/")
        except Exception as e:
            self.log.warning(f"[handin] legacy migration failed: {e}")

    @staticmethod
    def _safe_component(s: str, max_len: int = 80) -> str:
        s = (s or "").strip()
        s = re.sub(r'[<>:"/\\|?*]', "_", s)
        s = re.sub(r"\s+", " ", s).strip()
        s = s.rstrip(" .")
        if not s:
            s = "_"
        if len(s) > max_len:
            s = s[:max_len].rstrip(" .") or "_"
        return s

    # ----- task ops -----
    def list_active_tasks(self) -> List[HandinTask]:
        now = time.time()
        tasks = [t for t in self._tasks.values() if t.is_active(now)]
        tasks.sort(key=lambda x: x.deadline_ts)
        return tasks

    def list_active_tasks_by_group(self, group_id: int) -> List[HandinTask]:
        return [t for t in self.list_active_tasks() if int(t.group_id) == int(group_id)]


    def list_active_tasks_by_creator(self, creator_id: int) -> List[HandinTask]:
        """列出某个发起人创建的正在进行任务（跨群）。"""
        return [t for t in self.list_active_tasks() if int(t.creator_id) == int(creator_id)]

    # ===== 新增：列出任务（包含已截止/已结束/已取消）=====
    def list_tasks(self, include_closed: bool = True) -> List[HandinTask]:
        """列出任务。include_closed=True 时包含已截止/已结束/已取消的任务。"""
        tasks = list(self._tasks.values())
        if not include_closed:
            tasks = [t for t in tasks if not t.closed]
        # 近期优先：按截止时间倒序
        tasks.sort(key=lambda x: float(x.deadline_ts), reverse=True)
        return tasks

    def list_tasks_by_group(self, group_id: int, include_closed: bool = True) -> List[HandinTask]:
        """列出某群的任务（含已截止）。"""
        return [t for t in self.list_tasks(include_closed=include_closed) if int(t.group_id) == int(group_id)]

    def list_tasks_by_creator(self, creator_id: int, include_closed: bool = True) -> List[HandinTask]:
        """列出某个发起人创建的任务（跨群，含已截止）。"""
        return [t for t in self.list_tasks(include_closed=include_closed) if int(t.creator_id) == int(creator_id)]

    def list_submitted_files(self, task: HandinTask) -> List[Path]:
        """列出某任务已提交的文件（按修改时间倒序）。"""
        files_dir = self._task_files_dir(task.group_id, task.name)
        if not files_dir.exists():
            return []
        out = [p for p in files_dir.iterdir() if p.is_file() and not self._is_hash_index_file(p.name)]
        out.sort(key=lambda p: p.stat().st_mtime, reverse=True)
        return out

    @staticmethod
    def _is_hash_index_file(filename: str) -> bool:
        name = str(filename or "")
        return name == HANDIN_HASH_INDEX_FILENAME or name == f"{HANDIN_HASH_INDEX_FILENAME}.tmp"

    @staticmethod
    def file_sha256(path: Path, chunk_size: int = 1024 * 1024) -> str:
        h = hashlib.sha256()
        with Path(path).open("rb") as f:
            while True:
                b = f.read(chunk_size)
                if not b:
                    break
                h.update(b)
        return h.hexdigest()

    @staticmethod
    def _normalize_sha256(value: object) -> str:
        s = str(value or "").strip().lower()
        return s if re.fullmatch(r"[0-9a-f]{64}", s) else ""

    def _hash_index_path(self, files_dir: Path) -> Path:
        return Path(files_dir) / HANDIN_HASH_INDEX_FILENAME

    def _load_hash_index(self, files_dir: Path) -> Dict[str, dict]:
        path = self._hash_index_path(files_dir)
        try:
            obj = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            return {}
        raw = obj.get("files") if isinstance(obj, dict) else None
        if not isinstance(raw, dict):
            return {}
        out: Dict[str, dict] = {}
        for name, entry in raw.items():
            if not isinstance(entry, dict):
                continue
            fn = str(name or "").strip()
            sha = self._normalize_sha256(entry.get("sha256"))
            if not fn or self._is_hash_index_file(fn) or not sha:
                continue
            try:
                size = int(entry.get("size") or 0)
            except Exception:
                size = 0
            try:
                mtime_ns = int(entry.get("mtime_ns") or 0)
            except Exception:
                mtime_ns = 0
            out[fn] = {"sha256": sha, "size": size, "mtime_ns": mtime_ns}
        return out

    def _save_hash_index(self, files_dir: Path, entries: Dict[str, dict]) -> None:
        payload = {
            "version": 1,
            "updated_ts": int(time.time()),
            "files": entries,
        }
        path = self._hash_index_path(files_dir)
        tmp = path.with_suffix(path.suffix + ".tmp")
        tmp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        tmp.replace(path)

    def _refresh_hash_index(self, files_dir: Path) -> Dict[str, dict]:
        files_dir = Path(files_dir)
        entries = self._load_hash_index(files_dir)
        changed = False

        for name in list(entries.keys()):
            p = files_dir / name
            if not p.is_file():
                entries.pop(name, None)
                changed = True

        for p in files_dir.iterdir():
            if not p.is_file() or self._is_hash_index_file(p.name):
                continue
            try:
                st = p.stat()
            except Exception:
                continue
            old = entries.get(p.name) or {}
            old_sha = self._normalize_sha256(old.get("sha256"))
            if (
                old_sha
                and int(old.get("size") or -1) == int(st.st_size)
                and int(old.get("mtime_ns") or -1) == int(st.st_mtime_ns)
            ):
                continue
            entries[p.name] = {
                "sha256": self.file_sha256(p),
                "size": int(st.st_size),
                "mtime_ns": int(st.st_mtime_ns),
            }
            changed = True

        if changed:
            self._save_hash_index(files_dir, entries)
        return entries

    def zip_submissions(self, task: HandinTask, out_zip: Path) -> Tuple[bool, str, Optional[Path]]:
        """将某任务已提交文件全部打包为 zip。"""
        if getattr(task, "purged", False) or (not self._task_files_dir(task.group_id, task.name).exists()):
            return False, "该任务归档已超过保留期（最后一次 /handinget 后已清理），无法再导出。如需长期保留请及时备份。", None
        files = self.list_submitted_files(task)
        out_zip = Path(out_zip)
        out_zip.parent.mkdir(parents=True, exist_ok=True)
        try:
            with open_fast_zip(out_zip) as z:
                for p in files:
                    zip_write_path(z, p, arcname=p.name)
            return True, f"已打包 {len(files)} 个文件：{out_zip.name}", out_zip
        except Exception as e:
            return False, f"打包失败：{e}", None

    def create_task(
        self,
        group_id: int,
        creator_id: int,
        name: str,
        remind_ts_list: Optional[List[float]],
        deadline_ts: float,
        required_suffix: str = "",
    ) -> Tuple[bool, str]:
        """创建任务：提醒时间可为空或多个；最后一个时间一定是截止时间（由命令解析保证）。"""
        name = (name or "").strip()
        if not name or " " in name:
            return False, "任务名不合法：不能为空且不能包含空格。"
        if deadline_ts is None:
            return False, "时间格式不对：请用 月.日 时:分，例如 1.22 18:30（冒号中英文都行）。"
        suffix = ""
        if str(required_suffix or "").strip():
            suffix = normalize_required_suffix(required_suffix)
            if not suffix:
                return False, "文件后缀不支持。"

        rlist = []
        for x in (remind_ts_list or []):
            try:
                if x is None:
                    continue
                rlist.append(float(x))
            except Exception:
                continue
        # 去重并排序
        if rlist:
            rlist = sorted(set(rlist))

        dts = float(deadline_ts)
        if rlist and rlist[-1] >= dts:
            return False, "提醒时间必须早于截止时间。"

        # 同群同名且未截止的任务不允许重复创建
        for t in self._tasks.values():
            if int(t.group_id) == int(group_id) and t.name == name and t.is_active():
                return False, f"任务已存在：{name}（该群内同名任务尚未截止）"

        tid = f"{int(group_id)}:{name}:{int(time.time())}"
        task = HandinTask(
            task_id=tid,
            group_id=int(group_id),
            creator_id=int(creator_id),
            name=name,
            created_ts=time.time(),
            required_suffix=suffix,
            remind_ts_list=rlist,
            remind_sent_idx=0,
            deadline_ts=dts,
        )
        self._tasks[tid] = task
        # 创建目录
        self._task_files_dir(task.group_id, task.name)
        self._save()

        msg_lines = [f"创建提交任务成功：{name}"]
        if task.remind_ts_list:
            for i, ts in enumerate(task.remind_ts_list, 1):
                msg_lines.append(f"提醒{i}：{pretty_ts(ts)}")
        else:
            msg_lines.append("提醒：无")
        msg_lines.append(f"截止：{pretty_ts(task.deadline_ts)}")
        if task.required_suffix:
            msg_lines.append(f"限定格式：{required_suffix_display(task.required_suffix)}")
        return True, "\n".join(msg_lines)
    
    def cancel_task(self, task_id: str, by_user_id: int) -> Tuple[bool, str]:
        """取消任务：将任务标记为 closed/cancelled，停止后续提醒与截止推送。"""
        tid = str(task_id)
        t = self._tasks.get(tid)
        if not t:
            return False, "任务不存在。"
        if t.closed:
            return False, "任务已结束/已取消。"
        # 标记取消
        t.closed = True
        t.deadline_sent = True
        t.cancelled = True
        t.cancelled_ts = time.time()
        t.cancelled_by = int(by_user_id)
        self._save()
        return True, f"已取消任务「{t.name}」（群 {t.group_id}）。"

# ----- submissions -----
    def _unique_path(self, dst_dir: Path, filename: str) -> Path:
        filename = self._safe_component(filename, max_len=120)
        p = dst_dir / filename
        if not p.exists():
            return p
        stem = p.stem
        suf = p.suffix
        for i in range(2, 999):
            p2 = dst_dir / f"{stem}_{i}{suf}"
            if not p2.exists():
                return p2
        return dst_dir / f"{stem}_{int(time.time())}{suf}"

    @staticmethod
    def _normalize_download_url(url: str, filename: str) -> str:
        """QQ/FTN 下载链接有时会带空的 fname= 参数；补全它可显著提高可下载成功率。"""
        url = (url or "").strip()
        if not url:
            return ""
        try:
            sp = urllib.parse.urlsplit(url)
            pairs = urllib.parse.parse_qsl(sp.query, keep_blank_values=True)
            has_fname = False
            new_pairs = []
            for k, v in pairs:
                if k == "fname":
                    has_fname = True
                    if not v:
                        new_pairs.append((k, filename))
                    else:
                        new_pairs.append((k, v))
                else:
                    new_pairs.append((k, v))

            # 某些链接根本没有 fname，但实际也需要
            if (not has_fname) and ("ftn_handler" in sp.path):
                new_pairs.append(("fname", filename))

            new_q = urllib.parse.urlencode(new_pairs, doseq=True, encoding="utf-8", errors="strict")
            return urllib.parse.urlunsplit((sp.scheme, sp.netloc, sp.path, new_q, sp.fragment))
        except Exception:
            # 兜底：最简单的补全
            if url.endswith("fname="):
                return url + urllib.parse.quote(filename)
            return url

    @staticmethod
    def _pick_latest_temp_match(temp_dir: Path, *names: str) -> Optional[Path]:
        """在 NapCat temp 里按文件名/前缀兜底匹配，返回最新文件。"""
        if not temp_dir.exists() or not temp_dir.is_dir():
            return None

        patterns: list[str] = []
        for raw_name in names:
            nm = (raw_name or "").strip()
            if not nm:
                continue
            p = Path(nm)
            patterns.append(p.name)
            if p.suffix:
                patterns.append(f"{p.stem}*{p.suffix}")

        seen = set()
        hits: list[Path] = []
        for pat in patterns:
            if (not pat) or (pat in seen):
                continue
            seen.add(pat)
            try:
                for m in temp_dir.glob(pat):
                    if m.is_file():
                        hits.append(m)
            except Exception:
                continue

        if not hits:
            return None
        hits.sort(key=lambda p: p.stat().st_mtime, reverse=True)
        return hits[0]

    def download_to_inbox(
            self,
            user_id: int,
            fname: str,
            url: str,
            expected_size: Optional[int] = None,
            timeout: float = 180.0,
    ) -> Tuple[bool, str, Optional[Path]]:
        """保存私聊文件到 inbox。

        url 可能是：
        1) http/https 直链（可直接下载）
        2) NapCat 容器内本地路径（例如 /app/.config/QQ/NapCat/temp/xxx）
        3) file:///... 形式
        """
        # 兼容：fname 为空时兜底
        fname = (fname or "file").strip()

        raw = (url or "").strip()
        if not raw:
            return False, "文件缺少下载链接（url）。", None

        user_dir = self.inbox_dir / str(int(user_id))
        user_dir.mkdir(parents=True, exist_ok=True)

        # 生成唯一目标路径
        dst = self._unique_path(user_dir, fname)
        dst_part = dst.with_suffix(dst.suffix + ".part")

        # -------- 1) 处理 file:/// 路径 --------
        if raw.startswith("file:///"):
            # file:///C:/xxx 或 file:////app/... 都可能出现，统一转成本地路径字符串
            u = urlparse(raw)
            raw = unquote(u.path)  # Linux: /app/... ; Windows 可能是 /C:/...
            if os.name == "nt" and raw.startswith("/") and len(raw) >= 4 and raw[2] == ":":
                raw = raw[1:]  # /C:/xx -> C:/xx

        # -------- 2) 处理容器内本地缓存路径（以 / 开头）--------
        if raw.startswith("/"):
            try:
                temp_dir = Path(NAPCAT_TEMP_HOST_DIR)
                cdir = str(NAPCAT_TEMP_CONTAINER_DIR).rstrip("/")
                # 若路径在 NapCat temp 下，按映射关系找到宿主机对应文件
                if raw.startswith(cdir + "/") or raw == cdir:
                    rel = raw[len(cdir):].lstrip("/")
                    src = temp_dir / rel
                else:
                    # 兜底：按 basename 在 temp 目录找
                    src = temp_dir / Path(raw).name

                # NapCat 事件有时早于缓存落盘，给几秒等待并做一次模糊匹配兜底
                deadline = time.time() + 8.0
                while not (src.exists() and src.is_file()):
                    alt = self._pick_latest_temp_match(temp_dir, Path(raw).name, fname)
                    if alt is not None:
                        src = alt
                        break
                    if time.time() >= deadline:
                        break
                    time.sleep(0.4)

                if not src.exists() or not src.is_file():
                    if not temp_dir.exists():
                        return False, f"下载文件失败：NapCat 本地缓存目录不存在：{temp_dir}（请检查 NAPCAT_TEMP_HOST_DIR）", None
                    return False, f"下载文件失败：NapCat 本地缓存文件不存在：{src}", None

                # 大文件可能正在落盘：稍等，等 size 有明显增长/达到一定比例
                for _ in range(12):
                    if expected_size:
                        try:
                            exp = int(expected_size)
                            if exp > 0 and src.stat().st_size < max(32, exp // 10):
                                time.sleep(0.5)
                                continue
                        except Exception:
                            pass
                    break

                shutil.copy2(src, dst)
                size = dst.stat().st_size
                return True, f"已收到文件：{dst.name}（{size} bytes，本地缓存拷贝）", dst
            except Exception as e:
                return False, f"下载文件失败：本地缓存拷贝异常：{e}", None

        # -------- 3) http/https 下载 --------
        if raw.startswith("http://") or raw.startswith("https://"):
            try:
                req = urllib.request.Request(
                    raw,
                    headers={
                        "User-Agent": "Mozilla/5.0",
                    },
                )

                # 用 .part 写入，避免半截文件被当成成功
                if dst_part.exists():
                    try:
                        dst_part.unlink()
                    except Exception as e:
                        self.log.warning(f"[handin] remove stale partial file failed: path={dst_part} err={e}")

                with urllib.request.urlopen(req, timeout=float(timeout)) as resp, open(dst_part, "wb") as f:
                    downloaded = 0
                    chunk = 1024 * 1024  # 1MB
                    while True:
                        b = resp.read(chunk)
                        if not b:
                            break
                        f.write(b)
                        downloaded += len(b)

                        # 如果 expected_size 已知，且一直几乎不增长，也可以给时间让网络缓冲
                        # 这里不主动中断，只做稳妥写入

                # 下载完成后改名
                os.replace(dst_part, dst)

                size = dst.stat().st_size
                # 若 expected_size 存在且差距离谱，提示可能不完整
                if expected_size:
                    try:
                        exp = int(expected_size)
                        if exp > 0 and size < exp * 0.5:
                            return False, f"下载疑似不完整：期望约 {exp} bytes，实际 {size} bytes（可能链接失效/被拦截）", None
                    except Exception:
                        pass

                return True, f"已收到文件：{dst.name}（{size} bytes，网络下载）", dst
            except Exception as e:
                # 清理 part
                try:
                    if dst_part.exists():
                        dst_part.unlink()
                except Exception as e2:
                    self.log.warning(f"[handin] cleanup partial file after download error failed: path={dst_part} err={e2}")
                return False, f"下载文件失败：网络下载异常：{e}", None

        return False, f"不支持的下载来源：{raw}", None


    def move_inbox_to_task(self, inbox_path: Path, task: HandinTask, overwrite: bool = False) -> Tuple[bool, str, Optional[Path], str]:
        """将 inbox 临时文件移动到任务 files 目录。

        - 若目标存在同名文件且 overwrite=False：返回 code='EXISTS' 并不移动文件
        - overwrite=True：覆盖目标文件
        """
        if not inbox_path or not Path(inbox_path).exists():
            return False, "临时文件不存在（可能已过期/被清理）。", None, "MISSING"

        dst_dir = self._task_files_dir(task.group_id, task.name)
        dst_dir.mkdir(parents=True, exist_ok=True)
        dst = dst_dir / Path(inbox_path).name

        if dst.exists() and not overwrite:
            return False, f"任务「{task.name}」中已存在同名文件：{dst.name}", dst, "EXISTS"

        try:
            try:
                src_hash = self.file_sha256(inbox_path)
                entries = self._refresh_hash_index(dst_dir)
            except Exception as e:
                return False, f"文件校验失败：{e}", None, "ERR"

            for existing_name, entry in entries.items():
                if overwrite and existing_name == dst.name:
                    continue
                if self._normalize_sha256(entry.get("sha256")) == src_hash:
                    return (
                        False,
                        f"任务「{task.name}」中已存在内容完全相同的提交文件：{existing_name}，本次提交已终止。",
                        dst_dir / existing_name,
                        "DUPLICATE",
                    )

            if dst.exists() and overwrite:
                dst.unlink()
            Path(inbox_path).replace(dst)
            st = dst.stat()
            entries[dst.name] = {
                "sha256": src_hash,
                "size": int(st.st_size),
                "mtime_ns": int(st.st_mtime_ns),
            }
            self._save_hash_index(dst_dir, entries)
            return True, f"已归档到任务「{task.name}」：{dst.name}", dst, "OK"
        except Exception as e:
            return False, f"归档失败：{e}", None, "ERR"
    # ----- roster compare -----
    def compute_missing(self, task: HandinTask) -> Tuple[bool, str, List[Tuple[str, str]], Dict]:
        """返回 (ok, msg, missing_list, stats)."""
        roster = self._get_roster()
        if not roster:
            return False, f"读取班级名册失败：{ROSTER_XLSX_PATH}（文件不存在或格式不对）", [], {}

        submitted_ids: Set[str] = set()
        submitted_names: Set[str] = set()
        submitted_file_names: List[str] = []
        unknown_name_files: List[str] = []
        matched_name_files = 0

        roster_name_set = {str(nm or "").strip() for _, nm in roster if str(nm or "").strip()}
        roster_names = sorted(roster_name_set, key=lambda s: len(s), reverse=True)

        for p in self.list_submitted_files(task):
            # 统计所有已提交文件；仅跳过隐藏文件与临时分片
            if p.name.startswith("."):
                continue
            if p.suffix.lower() == ".part":
                continue
            submitted_file_names.append(p.name)

            sid = extract_student_id(p.name)
            if sid:
                submitted_ids.add(sid)

            nm = self.find_roster_name_in_filename(p.name, roster_names=roster_names)
            if not nm:
                # 兼容旧规则：先抽取姓名，再检查是否确实在名册中
                nm_guess = extract_name_from_filename(p.name)
                if nm_guess and (nm_guess in roster_name_set):
                    nm = nm_guess

            if nm:
                submitted_names.add(nm)
                matched_name_files += 1
            else:
                unknown_name_files.append(p.name)

        missing = []
        handed = 0
        for sid, nm in roster:
            if (sid and sid in submitted_ids) or (nm and nm in submitted_names):
                handed += 1
            else:
                missing.append((sid, nm))

        stats = {
            "roster_total": len(roster),
            "handed_in": handed,
            "missing": len(missing),
            "submitted_ids": len(submitted_ids),
            "submitted_names": len(submitted_names),
            "submitted_files_total": len(submitted_file_names),
            "recognized_name_files": matched_name_files,
            "recognized_name_ratio": (float(matched_name_files) / float(len(submitted_file_names))) if submitted_file_names else 1.0,
            "unknown_name_files": unknown_name_files,
            "submitted_file_names": submitted_file_names,
        }
        ratio = float(stats.get("recognized_name_ratio", 1.0))
        total_files = int(stats.get("submitted_files_total", 0))
        stats["use_submitted_list"] = bool(total_files > 0 and (matched_name_files <= 0 or ratio < 0.2))
        return True, "ok", missing, stats

    def format_missing_message(self, task: HandinTask, missing: List[Tuple[str, str]], stats: Dict, title: str) -> str:
        lines: List[str] = []
        lines.append(f"{title}\n任务：{task.name}\n群：{task.group_id}\n截止：{pretty_ts(task.deadline_ts)}")
        lines.append(f"已交/总人数：{stats.get('handed_in',0)}/{stats.get('roster_total',0)}；未交：{stats.get('missing',0)}")
        total_files = int(stats.get("submitted_files_total", 0) or 0)
        matched_name_files = int(stats.get("recognized_name_files", 0) or 0)
        ratio = float(stats.get("recognized_name_ratio", 0.0) or 0.0)
        if total_files > 0:
            lines.append(f"姓名识别文件占比：{matched_name_files}/{total_files}（{ratio * 100:.1f}%）")

        submitted_file_names = list(stats.get("submitted_file_names") or [])
        unknown_name_files = list(stats.get("unknown_name_files") or [])
        use_submitted_list = bool(stats.get("use_submitted_list"))

        # 文件名识别率太低时，未交名单准确性不足，改发已交文件列表
        if use_submitted_list:
            lines.append("⚠️ 姓名识别率过低（<20%）或未识别到名册姓名，改为发送已提交文件列表。")
            if not submitted_file_names:
                lines.append("当前没有已提交文件。")
                return "\n".join(lines)
            lines.append("已提交文件列表：")
            limit_files = 120
            for i, fn in enumerate(submitted_file_names[:limit_files], 1):
                lines.append(f"{i}. {fn}")
            if len(submitted_file_names) > limit_files:
                lines.append(f"...（共 {len(submitted_file_names)} 个，已截断显示前 {limit_files} 个）")
            return "\n".join(lines)

        if not missing:
            lines.append("✅ 全部已提交。")
        else:
            lines.append("未交名单：")
            limit = 120
            for i, (sid, nm) in enumerate(missing[:limit], 1):
                if nm:
                    lines.append(f"{i}. {nm}")
                else:
                    lines.append(f"{i}. （未知）")
            if len(missing) > limit:
                lines.append(f"...（共 {len(missing)} 人，已截断显示前 {limit} 人）")

        # 额外列出“已提交但未识别出姓名信息”的文件名
        if unknown_name_files:
            lines.append("")
            lines.append("未识别到姓名信息的已提交文件：")
            limit_unknown = 80
            for i, fn in enumerate(unknown_name_files[:limit_unknown], 1):
                lines.append(f"{i}. {fn}")
            if len(unknown_name_files) > limit_unknown:
                lines.append(f"...（共 {len(unknown_name_files)} 个，已截断显示前 {limit_unknown} 个）")
        return "\n".join(lines)

    # ----- scheduler -----
    async def scheduler_loop(self, api):
        """定时检查提醒/截止。"""
        while True:
            try:
                await asyncio.sleep(10)
                now = time.time()
                changed = False

                # 周期性清理（归档 + inbox）。默认每 1 小时最多跑一次。
                try:
                    if now - float(self._last_cleanup_ts or 0.0) >= 3600.0:
                        if self.cleanup_archives_and_inbox(now=now):
                            changed = True
                        self._last_cleanup_ts = now
                except Exception as e:
                    self.log.warning(f"[handin] periodic cleanup failed: err={e}")
                for t in list(self._tasks.values()):
                    if t.closed:
                        continue

                    # remind (0~N 次)
                    while t.remind_sent_idx < len(t.remind_ts_list) and now >= float(t.remind_ts_list[t.remind_sent_idx]):
                        idx = int(t.remind_sent_idx)
                        title = "📌 作业提交提醒"
                        if len(t.remind_ts_list) > 1:
                            title = f"📌 作业提交提醒（第 {idx+1}/{len(t.remind_ts_list)} 次）"
                        ok, msg, missing, stats = self.compute_missing(t)
                        if ok:
                            text = self.format_missing_message(t, missing, stats, title)
                        else:
                            text = title + "\n" + msg
                        await api.send_private_msg(t.creator_id, text)
                        t.remind_sent_idx = idx + 1
                        changed = True

                    # deadline
                    if (not t.deadline_sent) and now >= float(t.deadline_ts):
                        ok, msg, missing, stats = self.compute_missing(t)
                        if ok:
                            text = self.format_missing_message(t, missing, stats, "⏰ 作业截止提醒（已到截止时间）")
                        else:
                            text = "⏰ 作业截止提醒\n" + msg
                        await api.send_private_msg(t.creator_id, text)
                        t.deadline_sent = True
                        t.closed = True
                        changed = True

                if changed:
                    self._save()
            except Exception:
                self.log.exception("handin scheduler error")
